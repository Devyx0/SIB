"""
ingest_building_employee_assignment.py — Cleaning staff assignments from Master du ménage.

Source: '0 - Master du ménage des immeubles', tab 'Tous les immeubles'
Drive ID: 1hy-0mLsi4V9JMWxYBu-M-PQTXkyBF96ucpZiNkxLeeg

Building matching: reuses _match_building from parse_calendar_requests (returns
TABLE_FONDAMENTALE rows keyed by unique_building_name). Those names are then
looked up in the DB to get building_id.

Group-row detection: when DB building count > ménage row count for the matched
project, the ménage row covers multiple buildings — expand via Python keyword
filter on the project's sibling building addresses.

Run:
    uv run python -m drive_pipeline.ingestion.ingest_building_employee_assignment          # dry run
    uv run python -m drive_pipeline.ingestion.ingest_building_employee_assignment --write  # write
"""
from __future__ import annotations

import datetime
import io
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv

ROOT = Path(__file__).parent.parent.parent
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build as _gapi_build
from googleapiclient.http import MediaIoBaseDownload
import openpyxl
from psycopg2.extras import execute_values

from drive_pipeline.ingestion.db import get_conn
from drive_pipeline.ingestion.parse_calendar_requests import _match_building

_SCOPES   = ["https://www.googleapis.com/auth/drive"]
_DRIVE_ID = "1hy-0mLsi4V9JMWxYBu-M-PQTXkyBF96ucpZiNkxLeeg"
_TAB      = "Tous les immeubles"
_SOURCE   = "0 - Master du ménage des immeubles"

# Ménage sheet uses English project names; TABLE_FONDAMENTALE uses French for some.
_PROJ_ALIASES: dict[str, str] = {
    "the 48":                  "Le 48",
    "the claridge":            "Le Claridge",
    "place samuel de champlain": "Samuel-de-Champlain",
}

# Only Marie-Claude is first-name-only in this sheet; she's unique in the roster.
_FIRST_NAME_ONLY = {"marie-claude": 22, "marie claude": 22}


def _nrm(s: str) -> str:
    """Lowercase + strip accents + hyphens → spaces. Matches parse_calendar_requests._nrm."""
    s = unicodedata.normalize("NFD", s.lower())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s.replace("-", " ").strip()


def _build_emp_map(cur) -> dict[str, int]:
    """normalized-name → employee_id, including first-two-word shortcuts."""
    cur.execute("SELECT employee_id, nom_employe FROM employee")
    m: dict[str, int] = dict(_FIRST_NAME_ONLY)
    for eid, name in cur.fetchall():
        nrm = _nrm(name)
        m[nrm] = eid
        parts = nrm.split()
        if len(parts) >= 2:
            m[f"{parts[0]} {parts[1]}"] = eid  # "nathalie garneau simard" → "nathalie garneau"
    return m


def _parse_hours(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return None  # date stored in hours cell (data error, e.g. row 73 winter)
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().upper()
    if s in ("N/A", "NA", ""):
        return None
    return float(s.replace(",", "."))


def _extract_keywords(first_line: str) -> list[str]:
    """Street-name tokens from the first address line (accent-stripped ≥4 chars)."""
    keywords = []
    for segment in re.split(r"[/+]", first_line):
        seg = segment.strip()
        seg = re.sub(r"\b(?:Rue|Av|Boul|Ch|Rte|Avenue|Boulevard|Chemin|Route)\b", "", seg, flags=re.I)
        seg = re.sub(r"\b\d[\d\-àa\s]*", "", seg)
        seg = re.sub(r"\(.*?\)", "", seg)
        for word in seg.split():
            word = word.strip(" ,.")
            if len(word) >= 4:
                keywords.append(_nrm(word))
    return list(dict.fromkeys(keywords))


def _fetch_sheet() -> list[tuple[int, list]]:
    """Download sheet and return list of (sheet_row_1indexed, [col_values])."""
    creds = Credentials.from_authorized_user_file(str(ROOT / "token.json"), _SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    drive = _gapi_build("drive", "v3", credentials=creds)
    buf = io.BytesIO()
    req = drive.files().export_media(
        fileId=_DRIVE_ID,
        mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.seek(0)
    wb = openpyxl.load_workbook(buf, data_only=True)
    ws = wb[_TAB]
    rows = []
    for r in range(2, ws.max_row + 1):  # skip header row 1
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in vals):
            rows.append((r, vals))
    return rows


def run(dry_run: bool = True) -> None:
    conn = get_conn()
    cur  = conn.cursor()

    # ── Load DB context ────────────────────────────────────────────────────────
    emp_map = _build_emp_map(cur)

    # unique_building_name (normalized) → building_id
    cur.execute("SELECT building_id, adresse FROM building")
    bid_by_name: dict[str, int] = {_nrm(addr): bid for bid, addr in cur.fetchall()}

    # project_id → [(building_id, adresse_nrm), ...]
    cur.execute("SELECT building_id, project_id, adresse FROM building ORDER BY building_id")
    proj_buildings: dict[int, list[tuple[int, str]]] = defaultdict(list)
    bid_to_proj: dict[int, int] = {}
    for bid, pid, addr in cur.fetchall():
        proj_buildings[pid].append((bid, _nrm(addr)))
        bid_to_proj[bid] = pid

    # ── Fetch sheet ────────────────────────────────────────────────────────────
    menage_rows = _fetch_sheet()

    # ménage row count per normalized project name (rows with a real employee)
    proj_menage_count: dict[str, int] = defaultdict(int)
    for _, vals in menage_rows:
        proj_raw = vals[2]
        emp_raw  = vals[5]
        if proj_raw and emp_raw and str(emp_raw).strip().upper() not in ("N/A", "NA"):
            proj_menage_count[_nrm(str(proj_raw))] += 1

    # ── Process rows ───────────────────────────────────────────────────────────
    assignment_rows: dict[tuple[int, int], dict] = {}
    unmatched_employees: list[str] = []
    unmatched_buildings:  list[tuple[int, str]] = []
    seen_emp_flags: set[str] = set()

    for sheet_row, vals in menage_rows:
        proj_raw  = vals[2]
        addr_raw  = vals[4]
        emp_raw   = vals[5]
        hrs_sum   = _parse_hours(vals[7])
        hrs_win   = _parse_hours(vals[8])
        notes     = vals[9]

        emp_str = str(emp_raw).strip() if emp_raw is not None else ""
        if emp_str.upper() in ("N/A", "NA", ""):
            continue

        # ── Employee ───────────────────────────────────────────────────────────
        emp_key = _nrm(emp_str)
        emp_id  = emp_map.get(emp_key)
        if emp_id is None:
            if emp_str not in seen_emp_flags:
                unmatched_employees.append(f"  R{sheet_row}: {emp_str!r} — no match in employee table")
                seen_emp_flags.add(emp_str)
            continue

        # ── Address parse ──────────────────────────────────────────────────────
        addr_str   = str(addr_raw).strip() if isinstance(addr_raw, str) else ""
        first_line = addr_str.split("\n")[0].strip() if addr_str else ""

        civic, street = None, None
        m = re.match(r"^(.+?)\s*,\s*(\d[\d\-àa\.\s]*)", first_line)
        if m:
            street = m.group(1).strip()
            civic  = re.split(r"[\-àa\s]", m.group(2).strip())[0]
        else:
            m2 = re.match(r"^(.+?)\s+(\d+)\s*$", first_line.strip())
            if m2:
                street = m2.group(1).strip()
                civic  = m2.group(2)
            else:
                m3 = re.match(r"^(\d+)\s+\S+\s+(.+)", first_line)
                if m3:
                    civic  = m3.group(1)
                    street = m3.group(2).strip()

        proj_str       = str(proj_raw).strip() if proj_raw else None
        proj_for_match = _PROJ_ALIASES.get(_nrm(proj_str or "").replace(" ", " "), proj_str)
        # Alias lookup uses the normalized (accent-stripped) project name
        for alias_key, alias_val in _PROJ_ALIASES.items():
            if _nrm(proj_str or "") == _nrm(alias_key):
                proj_for_match = alias_val
                break
        tf_row = _match_building(street, civic, proj_for_match)

        if tf_row is None:
            unmatched_buildings.append((sheet_row, first_line or str(addr_raw)))
            continue

        bn = tf_row.get("unique_building_name", "")
        matched_bid = bid_by_name.get(_nrm(bn))
        if matched_bid is None:
            unmatched_buildings.append((sheet_row, first_line or str(addr_raw)))
            continue

        # ── Group-row expansion ────────────────────────────────────────────────
        proj_id      = bid_to_proj.get(matched_bid)
        db_count     = len(proj_buildings.get(proj_id, []))
        menage_count = proj_menage_count.get(_nrm(proj_str or ""), 1)

        if db_count > menage_count and proj_id is not None:
            keywords = _extract_keywords(first_line)
            if keywords:
                expanded = [
                    bid for bid, baddr_nrm in proj_buildings[proj_id]
                    if any(kw in baddr_nrm for kw in keywords)
                ]
                matched_bids = expanded if expanded else [matched_bid]
            else:
                matched_bids = [matched_bid]
        else:
            matched_bids = [matched_bid]

        # ── Collect ────────────────────────────────────────────────────────────
        for bid in matched_bids:
            key = (bid, emp_id)
            if key not in assignment_rows:
                assignment_rows[key] = {
                    "building_id":          bid,
                    "employee_id":          emp_id,
                    "weekly_hours_summer":  hrs_sum,
                    "weekly_hours_winter":  hrs_win,
                    "access_notes":         str(notes).strip() if notes else None,
                    "sheet_row":            sheet_row,
                }

    # ── Report ─────────────────────────────────────────────────────────────────
    print(f"\n=== building_employee_assignment {'DRY RUN' if dry_run else 'WRITE'} ===")
    print(f"\n  Rows to insert: {len(assignment_rows)}")

    if unmatched_employees:
        print(f"\n  UNMATCHED EMPLOYEES ({len(seen_emp_flags)}) — require client attention:")
        for line in unmatched_employees:
            print(line)

    if unmatched_buildings:
        print(f"\n  UNMATCHED BUILDINGS ({len(unmatched_buildings)}):")
        for r, addr in unmatched_buildings:
            print(f"  R{r}: {addr!r}")

    if dry_run:
        print("\n  Sample rows (building_id, employee_id, summer_h, winter_h):")
        for key, row in sorted(assignment_rows.items())[:25]:
            print(f"    bldg={row['building_id']:>3}  emp={row['employee_id']:>2}  "
                  f"sum={str(row['weekly_hours_summer']):>5}  win={str(row['weekly_hours_winter']):>5}")
        if len(assignment_rows) > 25:
            print(f"    … and {len(assignment_rows)-25} more")
        print("\n  DRY RUN — pass --write to commit.")
        conn.close()
        return

    # ── Write ──────────────────────────────────────────────────────────────────
    try:
        cur.execute("SELECT COALESCE(MAX(provenance_id), 0) FROM data_provenance")
        prov_next = cur.fetchone()[0] + 1

        rows_list = sorted(assignment_rows.values(), key=lambda r: (r["building_id"], r["employee_id"]))
        execute_values(
            cur,
            """INSERT INTO building_employee_assignment
               (building_id, employee_id, weekly_hours_summer, weekly_hours_winter, access_notes)
               VALUES %s ON CONFLICT (building_id, employee_id) DO NOTHING""",
            [(r["building_id"], r["employee_id"],
              r["weekly_hours_summer"], r["weekly_hours_winter"], r["access_notes"])
             for r in rows_list],
            page_size=200,
        )
        conn.commit()

        cur.execute("SELECT COUNT(*) FROM building_employee_assignment")
        total = cur.fetchone()[0]

        prov = []
        for row in rows_list:
            for col, val in [
                ("building_id",          str(row["building_id"])),
                ("employee_id",          str(row["employee_id"])),
                ("weekly_hours_summer",  str(row["weekly_hours_summer"]) if row["weekly_hours_summer"] is not None else None),
                ("weekly_hours_winter",  str(row["weekly_hours_winter"]) if row["weekly_hours_winter"] is not None else None),
                ("access_notes",         row["access_notes"]),
            ]:
                prov.append((
                    prov_next,
                    "building_employee_assignment",
                    row["building_id"],
                    col,
                    "google_sheets", _DRIVE_ID, _SOURCE, _TAB,
                    row["sheet_row"], None,
                    val, val, None,
                ))
                prov_next += 1

        execute_values(
            cur,
            """INSERT INTO data_provenance
               (provenance_id, target_table, target_pk, target_column,
                source_type, source_drive_id, source_file_name, source_tab,
                source_row, source_column, raw_value, ingested_value, ingested_at)
               VALUES %s""",
            prov, page_size=500,
        )
        conn.commit()

        print(f"\n  building_employee_assignment: {total} total rows")
        print(f"  Provenance rows written: {len(prov)}")

    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    run(dry_run="--write" not in sys.argv)
