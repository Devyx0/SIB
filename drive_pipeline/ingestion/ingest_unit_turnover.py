"""
ingest_unit_turnover.py — Write unit move-in/move-out history from Arrivées et Départs sheet.

Source: Google Sheets ID _SHEET_ID
  NEW tab 'Departs Arrivees complétées': 871 rows, double header (row 1 stale, row 2 real),
    data from row 3. Unit ID in col B (idx 1); dates in cols E/F/H/S/T.
    Tenant names: idx 85 (incoming), idx 89 (outgoing).
  OLD tab 'Départs  Arrivées complétés': 544 rows, single header, col B = "Adresse"
    (inconsistent unit identifier — parsed best-effort).

Writes to: unit_turnover + data_provenance

Run:
    uv run python -m drive_pipeline.ingestion.ingest_unit_turnover          # dry run
    uv run python -m drive_pipeline.ingestion.ingest_unit_turnover --write  # write
"""
from __future__ import annotations

import datetime
import io
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from dotenv import load_dotenv

ROOT = Path(__file__).parent.parent.parent
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build as _gapi_build
from googleapiclient.http import MediaIoBaseDownload

from drive_pipeline.ingestion.db import get_conn
from drive_pipeline.ingestion.parse_calendar_requests import _match_building

_SCOPES   = ["https://www.googleapis.com/auth/drive"]
_SHEET_ID = "1jTTmL1Kfh3J0bTrrQnLW-kYQc9V0zrdbEtmn2nUixSM"
_SOURCE_FILE = "Tableau_des_Arrivees_et_Departs"

# Column indices in NEW tab row 2 (0-based, confirmed by probe)
_C_UNIT    = 1   # "Numero unique d'unite"
_C_END     = 4   # "DÉPART\nDate de fin de bail"
_C_OUT_DT  = 5   # "DÉPART\nDate de départ confirmée"
_C_CESSION = 7   # "DÉPART\nSi cession : Date de départ souhaitée"
_C_START   = 18  # "ARRIVÉE\nDate début de bail"
_C_IN_DT   = 19  # "ARRIVÉE\nDate d'emménagement confirmée"
_C_IN_NAME = 85  # "Locataires Futurs\nNoms"
_C_OUT_NAME = 89 # "Locataires Actuels\nNoms"


# ── helpers ──────────────────────────────────────────────────────────────────

def _norm_key(name: str) -> str:
    s = unicodedata.normalize("NFD", str(name))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s.lower())


def _clean(val) -> str | None:
    if val is None:
        return None
    s = str(val).strip()
    return None if s in ("", "-", "N/A", "n/a", "#N/A", "none") else s


def _parse_date(val) -> datetime.date | None:
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        return val.date()
    if isinstance(val, datetime.date):
        return val
    s = str(val).strip()
    if not s or s in ("-", "N/A"):
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _cell(row: tuple, idx: int):
    return row[idx].value if idx < len(row) else None


# ── Drive ────────────────────────────────────────────────────────────────────

def _get_drive():
    creds = Credentials.from_authorized_user_file(str(ROOT / "token.json"), _SCOPES)
    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
    return _gapi_build("drive", "v3", credentials=creds)


def _export_sheet(drive, file_id: str) -> bytes:
    buf = io.BytesIO()
    req = drive.files().export_media(
        fileId=file_id,
        mimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    return buf.getvalue()


# ── Unit resolution ───────────────────────────────────────────────────────────

def _resolve_unit(unit_raw: str, bldg_by_addr: dict,
                  unit_by_bldg_apt: dict) -> int | None:
    """
    Parse '{civic}-{apt} {street}' and resolve to unit_id.
    Strips notes in parens: '51-9 Langevin (CESSION...)' → '51-9 Langevin'.
    Apt may be alphanumeric: '460A', '4B'. Non-numeric words (Garage) give no match.
    TF match first; scan ALL buildings when TF gives wrong building (shared-civic clusters).
    """
    s = unit_raw.split("(")[0].strip().rstrip(",;")
    m = re.match(r"^(\d+)-([A-Za-z0-9]+)\s+(.+)$", s)
    if not m:
        return None
    civic, apt, street = m.group(1), m.group(2), m.group(3).strip()
    if not re.search(r"\d", apt):  # reject Garage, Loft, etc.
        return None

    # TF-based match (handles abbreviations like St- vs Saint-)
    tf_row = _match_building(street, civic, None)
    if tf_row:
        bid = bldg_by_addr.get(tf_row.get("unique_building_name"))
        if bid:
            uid = unit_by_bldg_apt.get((bid, apt))
            if uid:
                return uid
            # TF gave a building but apt absent — fall through to scan all buildings
            # (shared-civic clusters: 6 Chamonix at 1175, apt 197 needs bldg 38 not 33)

    # Scan ALL buildings with matching civic+street
    c_n = _norm_key(civic)
    s_words = [w for w in _norm_key(street).split() if len(w) >= 4]
    if not s_words:
        return None
    for addr, bid in bldg_by_addr.items():
        a_n = _norm_key(addr)
        if c_n in a_n and any(sw in a_n for sw in s_words):
            uid = unit_by_bldg_apt.get((bid, apt))
            if uid:
                return uid
    return None


# ── DB ────────────────────────────────────────────────────────────────────────

def _load_db(conn):
    cur = conn.cursor()

    cur.execute("SELECT adresse, building_id FROM building")
    bldg_by_addr = {r[0]: r[1] for r in cur.fetchall() if r[0]}

    cur.execute(
        "SELECT u.unit_id, u.building_id, u.numero_appartement FROM unit u"
    )
    unit_by_bldg_apt: dict[tuple, int] = {}
    for uid, bid, apt in cur.fetchall():
        if apt:
            unit_by_bldg_apt[(bid, str(apt).strip())] = uid

    cur.execute("SELECT tenant_id, nom_complet FROM tenant")
    tenant_by_norm = {_norm_key(str(r[1])): r[0] for r in cur.fetchall() if r[1]}

    cur.execute("SELECT COALESCE(MAX(turnover_id), 0) FROM unit_turnover")
    max_tid = cur.fetchone()[0]

    cur.execute("SELECT COALESCE(MAX(provenance_id), 0) FROM data_provenance")
    max_prov = cur.fetchone()[0]

    return bldg_by_addr, unit_by_bldg_apt, tenant_by_norm, max_tid, max_prov


# ── Parse tabs ────────────────────────────────────────────────────────────────

def _parse_new_tab(ws, bldg_by_addr: dict, unit_by_bldg_apt: dict,
                   tenant_by_norm: dict):
    """Parse 'Departs Arrivees complétées' (double header; data from row 3)."""
    rows_out = []
    skipped_stat = []
    unresolved = []

    for row in ws.iter_rows(min_row=3):
        unit_raw = _clean(_cell(row, _C_UNIT))
        if unit_raw is None:
            continue
        if unit_raw.upper().startswith("STAT"):
            skipped_stat.append(unit_raw)
            continue

        unit_id = _resolve_unit(unit_raw, bldg_by_addr, unit_by_bldg_apt)
        if unit_id is None:
            unresolved.append(unit_raw)

        inc_name = _clean(_cell(row, _C_IN_NAME))
        out_name = _clean(_cell(row, _C_OUT_NAME))
        inc_id = tenant_by_norm.get(_norm_key(inc_name)) if inc_name else None
        out_id = tenant_by_norm.get(_norm_key(out_name)) if out_name else None

        rows_out.append({
            "unit_raw":           unit_raw,
            "unit_id":            unit_id,
            "lease_end_date":     _parse_date(_cell(row, _C_END)),
            "move_out_confirmed": _parse_date(_cell(row, _C_OUT_DT)),
            "assignment_desired": _parse_date(_cell(row, _C_CESSION)),
            "lease_start_date":   _parse_date(_cell(row, _C_START)),
            "move_in_confirmed":  _parse_date(_cell(row, _C_IN_DT)),
            "incoming_tenant_id": inc_id,
            "outgoing_tenant_id": out_id,
            "incoming_name_raw":  inc_name,
            "outgoing_name_raw":  out_name,
            "source_tab":         "Departs Arrivees complétées",
        })
    return rows_out, skipped_stat, unresolved


def _parse_old_tab(ws, bldg_by_addr: dict, unit_by_bldg_apt: dict):
    """Parse 'Départs  Arrivées complétés' (single header; data from row 2).
    Col B = Adresse (unit identifier, inconsistent); Col D = Date emménagement.
    """
    rows_out = []
    unresolved = []

    for row in ws.iter_rows(min_row=2):
        unit_raw = _clean(_cell(row, 1))  # col B
        if unit_raw is None:
            continue
        move_in = _parse_date(_cell(row, 3))  # col D

        unit_id = _resolve_unit(unit_raw, bldg_by_addr, unit_by_bldg_apt)
        if unit_id is None:
            unresolved.append(unit_raw)
            continue

        rows_out.append({
            "unit_raw":           unit_raw,
            "unit_id":            unit_id,
            "lease_end_date":     None,
            "move_out_confirmed": None,
            "assignment_desired": None,
            "lease_start_date":   None,
            "move_in_confirmed":  move_in,
            "incoming_tenant_id": None,
            "outgoing_tenant_id": None,
            "incoming_name_raw":  None,
            "outgoing_name_raw":  None,
            "source_tab":         "Départs  Arrivées complétés",
        })
    return rows_out, unresolved


# ── Main ──────────────────────────────────────────────────────────────────────

def run(dry_run: bool = True) -> None:
    conn = get_conn()
    bldg_by_addr, unit_by_bldg_apt, tenant_by_norm, max_tid, max_prov = _load_db(conn)

    print(f"\n=== ingest_unit_turnover {'DRY RUN' if dry_run else 'WRITE'} ===")

    drive = _get_drive()
    print("Downloading Google Sheet…")
    xlsx_bytes = _export_sheet(drive, _SHEET_ID)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)

    tab_names = wb.sheetnames
    print(f"Tabs: {tab_names}")

    # Find NEW and OLD tabs by partial name match
    new_name = next((n for n in tab_names if "Departs Arrivees" in n), None)
    old_name = next((n for n in tab_names if "Départs" in n and "complétés" in n
                     and "DÉPARTS ET" not in n), None)
    if not new_name:
        print("ERROR: could not find NEW tab (Departs Arrivees…)")
        return

    print(f"\nNEW tab: {new_name!r}")
    ws_new = wb[new_name]
    new_rows, skipped_stat, new_unresolved = _parse_new_tab(
        ws_new, bldg_by_addr, unit_by_bldg_apt, tenant_by_norm
    )
    new_writable = [r for r in new_rows if r["unit_id"] is not None]

    print(f"OLD tab: {old_name!r}")
    old_rows_raw, old_unresolved = [], []
    old_total = 0
    if old_name:
        ws_old = wb[old_name]
        old_rows_raw, old_unresolved = _parse_old_tab(ws_old, bldg_by_addr, unit_by_bldg_apt)
        old_total = sum(1 for _ in ws_old.iter_rows(min_row=2)) - 1

    # Tenancy_record backfill estimate
    unit_ids_with_start = [r["unit_id"] for r in new_writable if r["lease_start_date"]]
    tr_backfill = 0
    if unit_ids_with_start:
        cur = conn.cursor()
        cur.execute(
            """
            SELECT COUNT(*) FROM tenancy_record tr
            JOIN lease l ON l.lease_id = tr.lease_id
            WHERE tr.start_date IS NULL
              AND l.unit_id = ANY(%s)
            """,
            (unit_ids_with_start,),
        )
        tr_backfill = cur.fetchone()[0]
    conn.close()

    # ── Dry run report ────────────────────────────────────────────────────────
    print(f"\n{'='*60}")
    print(f"NEW tab parsed rows:         {len(new_rows)}")
    print(f"  STAT_* parking rows skipped: {len(skipped_stat)}")
    print(f"  Unit resolved:               {len(new_writable)}")
    print(f"  Unit NOT resolved:           {len(new_unresolved)}")
    print(f"  Incoming tenant FK found:    {sum(1 for r in new_writable if r['incoming_tenant_id'])}")
    print(f"  Outgoing tenant FK found:    {sum(1 for r in new_writable if r['outgoing_tenant_id'])}")
    print()
    print(f"OLD tab rows total: {old_total}")
    print(f"  Unit resolved (writable):  {len(old_rows_raw)}")
    print(f"  Unit NOT resolved (skip):  {len(old_unresolved)}")
    print()
    print(f"tenancy_record backfill estimate:")
    print(f"  NEW rows with lease_start_date: {len(unit_ids_with_start)}")
    print(f"  tenancy_records with NULL start_date for those units: {tr_backfill}")

    if new_unresolved[:15]:
        print(f"\nNEW tab unresolved unit_ids (first 15 of {len(new_unresolved)}):")
        for u in new_unresolved[:15]:
            print(f"  - {u!r}")

    all_writable = new_writable + old_rows_raw
    print(f"\nTotal writable rows: {len(all_writable)}")

    if dry_run:
        print(f"\nDRY RUN — pass --write to commit.")
        return

    # ── Write ─────────────────────────────────────────────────────────────────
    conn = get_conn()
    try:
        cur = conn.cursor()
        cur.execute("SELECT COALESCE(MAX(provenance_id), 0) FROM data_provenance")
        prov_seq = cur.fetchone()[0]
        cur.execute("SELECT COALESCE(MAX(turnover_id), 0) FROM unit_turnover")
        tid_seq = cur.fetchone()[0]

        inserted = 0
        for r in all_writable:
            tid_seq += 1
            cur.execute(
                """INSERT INTO unit_turnover
                   (turnover_id, unit_id,
                    lease_end_date, move_out_confirmed_date, assignment_desired_date,
                    lease_start_date, move_in_confirmed_date,
                    incoming_tenant_id, outgoing_tenant_id,
                    incoming_tenant_name_raw, outgoing_tenant_name_raw)
                   VALUES (%s,%s, %s,%s,%s, %s,%s, %s,%s, %s,%s)""",
                (tid_seq, r["unit_id"],
                 r["lease_end_date"], r["move_out_confirmed"], r["assignment_desired"],
                 r["lease_start_date"], r["move_in_confirmed"],
                 r["incoming_tenant_id"], r["outgoing_tenant_id"],
                 r["incoming_name_raw"], r["outgoing_name_raw"]),
            )
            inserted += 1
            for col, val in [
                ("lease_end_date",          r["lease_end_date"]),
                ("move_out_confirmed_date", r["move_out_confirmed"]),
                ("assignment_desired_date", r["assignment_desired"]),
                ("lease_start_date",        r["lease_start_date"]),
                ("move_in_confirmed_date",  r["move_in_confirmed"]),
            ]:
                if val is None:
                    continue
                prov_seq += 1
                cur.execute(
                    """INSERT INTO data_provenance
                       (provenance_id, target_table, target_pk, target_column,
                        source_type, source_drive_id, source_file_name,
                        source_tab, source_row, source_column, raw_value, ingested_value)
                       VALUES (%s,'unit_turnover',%s,%s,'google_sheet',%s,%s,%s,NULL,NULL,%s,%s)""",
                    (prov_seq, str(tid_seq), col,
                     _SHEET_ID, _SOURCE_FILE, r["source_tab"],
                     str(val), str(val)),
                )

        conn.commit()
        print(f"\nCommitted: {inserted} unit_turnover rows, {prov_seq - max_prov} provenance rows")

    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    run(dry_run="--write" not in sys.argv)
