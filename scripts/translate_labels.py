"""
translate_labels.py

Translates column A (field labels) of all Fiches Techniques .xlsx files
from French to English via label_glossary.py. Column B (values) untouched.

Dry run (default):
    uv run translate_labels.py

Write back to Drive (after dry-run approval):
    uv run translate_labels.py --write

Formatting note: write-back uses openpyxl.load_workbook() directly — it loads
the existing workbook structure, modifies only the specific column-A cells that
need translating, then re-saves. Column B, all styles, borders, merged cells,
column widths, row heights, and formulas are fully preserved. This is NOT a
pandas to_excel() rewrite (which would discard formatting).
"""
import argparse
import io
import os
import sys
import threading
import concurrent.futures
from googleapiclient.discovery import build
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
import pandas as pd
import openpyxl
from drive_pipeline.fiches_techniques.links import FICHES
from drive_pipeline.fiches_techniques.label_glossary import EXCLUDED_LABELS, VARIANT_TO_CANONICAL, GLOSSARY

SCOPES = ["https://www.googleapis.com/auth/drive"]
TOKEN_FILE  = os.path.join(os.path.dirname(__file__), "token.json")

FICHES_SUBFOLDERS = {"root", "01-Old Technical Sheets", "00-Archives"}
EXCLUDED_NAMES = {
    "0 - EMERGENCY CONTACTS AND PROCEDURES.xlsx",
    "Untitled.xlsx",
    "Test Copy of St-Jean Street, Québec (636 640 646).xlsx",
}
DRY_RUN_REPORT   = os.path.join(os.path.dirname(__file__), "dry_run_report.py")
TRANSLATION_LOG  = os.path.join(os.path.dirname(__file__), "label_translation_log.py")

MAX_WORKERS = 8

_EXCLUDED_SET = set(EXCLUDED_LABELS)
_tls = threading.local()


# ── auth ──────────────────────────────────────────────────────────────────────

def get_service():
    if not hasattr(_tls, "service"):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())
        _tls.service = build("drive", "v3", credentials=creds)
    return _tls.service


def download_bytes(service, file_id):
    buf = io.BytesIO()
    req = service.files().get_media(fileId=file_id)
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.seek(0)
    return buf


# ── translation lookup ────────────────────────────────────────────────────────

def translate_label(raw):
    """
    Returns (result_label, status).
    status: 'excluded' | 'translated' | 'no-change' | 'unrecognized'
    """
    if raw in _EXCLUDED_SET:
        return raw, "excluded"
    canonical = VARIANT_TO_CANONICAL.get(raw)
    if canonical is None:
        return raw, "unrecognized"
    english = GLOSSARY.get(canonical)
    if english is None:
        return raw, "unrecognized"
    if english == raw:
        return raw, "no-change"
    return english, "translated"


# ── file scope ────────────────────────────────────────────────────────────────

def get_targets():
    return [
        f for f in FICHES
        if f["subfolder"] in FICHES_SUBFOLDERS
        and f["english_name"] not in EXCLUDED_NAMES
        and f["english_name"].endswith(".xlsx")
    ]


# ── per-file workers ──────────────────────────────────────────────────────────

def process_dry(fiche):
    service = get_service()
    file_id = fiche["english_drive_id"]
    try:
        buf = download_bytes(service, file_id)
        df  = pd.read_excel(buf, header=None, engine="openpyxl")
    except Exception as exc:
        return {"fiche": fiche, "error": str(exc),
                "changes": [], "excluded": [], "unrecognized": []}

    changes      = []
    excluded     = []
    unrecognized = []

    for i in range(1, df.shape[0]):   # skip row 0 (civic address)
        cell = df.iat[i, 0]
        if not isinstance(cell, str):
            continue
        label = cell.strip()
        if not label:
            continue
        english, status = translate_label(label)
        if status == "excluded":
            excluded.append({"row": i, "label": label})
        elif status == "unrecognized":
            unrecognized.append({"row": i, "label": label})
        elif status == "translated":
            changes.append({"row": i, "old_label": label, "new_label": english})
        # "no-change" silently skipped — nothing to do

    return {"fiche": fiche, "error": None,
            "changes": changes, "excluded": excluded, "unrecognized": unrecognized}


def process_write(fiche):
    service = get_service()
    file_id = fiche["english_drive_id"]
    try:
        buf = download_bytes(service, file_id)
        wb  = openpyxl.load_workbook(buf)   # preserves formatting, merged cells, formulas
    except Exception as exc:
        return {"fiche": fiche, "error": str(exc), "changes": []}

    ws      = wb.active
    changes = []

    for row_idx, row in enumerate(ws.iter_rows()):
        if row_idx == 0:
            continue   # skip civic address row
        cell_a = row[0]
        val = cell_a.value
        if not isinstance(val, str):
            continue
        label = val.strip()
        if not label:
            continue
        english, status = translate_label(label)
        if status == "translated":
            cell_a.value = english   # only this assignment touches the file
            changes.append({"row_index": row_idx, "old_label": label, "new_label": english})

    if changes:
        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        media = MediaIoBaseUpload(
            out_buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resumable=False,
        )
        service.files().update(fileId=file_id, media_body=media).execute()

    return {"fiche": fiche, "error": None, "changes": changes}


# ── output files ──────────────────────────────────────────────────────────────

def write_dry_report(results):
    lines = ["DRY_RUN = [\n"]
    for r in results:
        f = r["fiche"]
        lines.append(
            f"    {{'english_drive_id': {f['english_drive_id']!r}, "
            f"'english_name': {f['english_name']!r}, "
            f"'subfolder': {f['subfolder']!r}, "
            f"'error': {r.get('error')!r}, "
            f"'changes': {r['changes']!r}, "
            f"'excluded': {r.get('excluded', [])!r}, "
            f"'unrecognized': {r.get('unrecognized', [])!r}}},\n"
        )
    lines.append("]\n")
    with open(DRY_RUN_REPORT, "w", encoding="utf-8") as fh:
        fh.writelines(lines)


def write_translation_log(results):
    lines = ["TRANSLATION_LOG = [\n"]
    for r in results:
        if not r.get("changes"):
            continue
        f = r["fiche"]
        lines.append(
            f"    {{'english_drive_id': {f['english_drive_id']!r}, "
            f"'english_name': {f['english_name']!r}, "
            f"'subfolder': {f['subfolder']!r}, "
            f"'changes': {r['changes']!r}}},\n"
        )
    lines.append("]\n")
    with open(TRANSLATION_LOG, "w", encoding="utf-8") as fh:
        fh.writelines(lines)


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true",
                        help="Write translated labels back to Drive (default: dry run only)")
    args = parser.parse_args()

    targets = get_targets()
    mode    = "WRITE" if args.write else "DRY RUN"
    print(f"{mode}: {len(targets)} files")

    worker  = process_write if args.write else process_dry
    results = [None] * len(targets)
    done    = 0

    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        future_to_idx = {pool.submit(worker, t): i for i, t in enumerate(targets)}
        for future in concurrent.futures.as_completed(future_to_idx):
            idx = future_to_idx[future]
            results[idx] = future.result()
            done += 1
            if done % 50 == 0:
                print(f"  {done}/{len(targets)} done...")

    # ── collate ──
    errors       = [r for r in results if r.get("error")]
    no_changes   = [r for r in results if not r.get("error") and not r["changes"]]
    with_changes = [r for r in results if not r.get("error") and r["changes"]]
    total_cells  = sum(len(r["changes"]) for r in with_changes)

    all_unrecognized = [
        (r["fiche"]["english_name"], u)
        for r in results if not r.get("error")
        for u in r.get("unrecognized", [])
    ]
    all_excluded = [
        (r["fiche"]["english_name"], e)
        for r in results if not r.get("error")
        for e in r.get("excluded", [])
    ]

    # ── report ──
    print(f"\n{'='*60}")
    print(f"Files processed:      {len(results)}")
    print(f"Files with changes:   {len(with_changes)}")
    print(f"Files with 0 changes: {len(no_changes)}")
    print(f"Files errored:        {len(errors)}")
    print(f"Total labels changed: {total_cells}")
    print(f"Unrecognized labels:  {len(all_unrecognized)}")
    print(f"Excluded labels hit:  {len(all_excluded)}")

    if errors:
        print(f"\nERRORS:")
        for r in errors:
            print(f"  {r['fiche']['english_name']!r}: {r['error']}")

    if no_changes:
        print(f"\nZERO-CHANGE FILES ({len(no_changes)}):")
        for r in no_changes:
            print(f"  [{r['fiche']['subfolder']}] {r['fiche']['english_name']!r}")

    if all_unrecognized:
        print(f"\nUNRECOGNIZED LABELS (glossary gaps — NOT translated):")
        for fname, u in all_unrecognized:
            print(f"  {fname!r}  row={u['row']}  {u['label']!r}")

    if all_excluded:
        print(f"\nEXCLUDED LABELS HIT IN SCOPE FILES (left as-is):")
        for fname, e in all_excluded:
            print(f"  {fname!r}  row={e['row']}  {e['label']!r}")

    if not args.write:
        write_dry_report(results)
        print(f"\nFull detail → dry_run_report.py")
        print("Run with --write to apply.")
    else:
        write_translation_log(results)
        print(f"\nAudit trail → label_translation_log.py")


if __name__ == "__main__":
    main()
