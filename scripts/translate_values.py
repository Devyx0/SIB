"""
translate_values.py

Translates column B (French values) into English and writes to column C.
One Anthropic API call per file, all non-empty column B cells batched.

Dry run (default — local copies + report only, no Drive writes):
    uv run translate_values.py

Write to Drive:
    uv run translate_values.py --write
"""
import argparse
import io
import json
import math
import os
import re
import sys
import threading
import time
import concurrent.futures
from pathlib import Path

from googleapiclient.discovery import build
from google.oauth2.credentials import Credentials
from google.auth.transport.requests import Request
from googleapiclient.http import MediaIoBaseDownload, MediaIoBaseUpload
import openpyxl
import anthropic

from dotenv import load_dotenv
from drive_pipeline.fiches_techniques.links import FICHES

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"))

SCOPES     = ["https://www.googleapis.com/auth/drive"]
TOKEN_FILE = os.path.join(os.path.dirname(__file__), "token.json")

FICHES_SUBFOLDERS = {"root", "01-Old Technical Sheets", "00-Archives"}
EXCLUDED_NAMES = {
    "0 - EMERGENCY CONTACTS AND PROCEDURES.xlsx",
    "Untitled.xlsx",
    "Test Copy of St-Jean Street, Québec (636 640 646).xlsx",
}

LOCAL_DIR        = Path(os.path.dirname(__file__)) / "dryrun_col_c"
VALUE_DRY_REPORT = os.path.join(os.path.dirname(__file__), "value_dry_report.py")
VALUE_TRANS_LOG  = os.path.join(os.path.dirname(__file__), "value_translation_log.py")

MAX_WORKERS = 8
MODEL       = "claude-haiku-4-5-20251001"

_CAPS_RE = re.compile(r'\b([A-Z]{2,4})\b')

_tls = threading.local()

if not os.environ.get("ANTHROPIC_API_KEY"):
    sys.exit("Error: ANTHROPIC_API_KEY environment variable not set")

_extra_headers = {}
_wsid = os.environ.get("ANTHROPIC_WORKSPACE_ID", "")
if _wsid:
    _extra_headers["anthropic-workspace-id"] = _wsid
_client = anthropic.Anthropic(default_headers=_extra_headers if _extra_headers else None)


SYSTEM_PROMPT = """\
You are a precise French → English translation engine for property management records.

Apply these rules in strict priority order:

PRESERVE exactly as-is (status: "preserved_code_or_name"):
- Phone numbers, alarm codes, access codes, door codes, PIN codes
- Proper nouns: street names, company/vendor names (Lanvac, Bell, SIB, ILO,
  Viessmann, ESSA, MRA, Blais, Paradox, Akuvox, Médéco, Logisco, and any other
  recognizable proper noun or brand)
- Client/account reference codes (e.g. #Client A57921)
- Measurements and dimensions — preserve numbers and units exactly
- Short ALL-CAPS abbreviations of 2–4 letters (AV, AR, CE, SS, SB, PB, etc.)
  — these are physical labels on real keys, panels, or equipment; never expand or
  translate them
- ANY acronym or abbreviation you cannot confidently expand to a full English phrase
- Login credentials (usernames, passwords) — even if already in English
- Numeric-only values, years

TRANSLATE the prose (status: "translated"):
- Location descriptions, procedural instructions, descriptive content
- For MIXED cells: translate only the prose portions; leave all codes/names/
  numbers/abbreviations verbatim in place

NO-CHANGE (status: "no_change"):
- Cell is already fully in English — return original unchanged

LOW CONFIDENCE (status: "preserved_low_confidence"):
- When you are not fully confident the translation is operationally safe
- Return the original French text UNCHANGED as "translation"
- Never guess — wrong translations of security codes or location instructions
  can cause operational harm

OUTPUT: Return ONLY a valid JSON array. No markdown fences, no preamble.
One object per input cell, in the same order as input:
[
  {"row_idx": <int>, "translation": "<text>",
   "status": "translated" | "preserved_code_or_name" | "preserved_low_confidence" | "no_change"},
  ...
]"""


# ── auth ──────────────────────────────────────────────────────────────────────

def get_service():
    if not hasattr(_tls, "service"):
        creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())
        _tls.service = build("drive", "v3", credentials=creds)
    return _tls.service


def download_bytes(service, file_id):
    buf = io.BytesIO()
    req = service.files().get_media(fileId=file_id)
    dl  = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    buf.seek(0)
    return buf


# ── cell extraction ───────────────────────────────────────────────────────────

def is_empty(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str) and not v.strip():
        return True
    return False


def collect_cells(ws):
    """Return non-empty col-B cells (row_idx 0-based, skip row 0 = address)."""
    cells = []
    for row_idx, row in enumerate(ws.iter_rows()):
        if row_idx == 0:
            continue
        col_b = row[1] if len(row) > 1 else None
        if col_b is None or is_empty(col_b.value):
            continue
        col_a = row[0] if row else None
        label = col_a.value.strip() if col_a and isinstance(col_a.value, str) else ""
        cells.append({
            "row_idx": row_idx,
            "label":   label,
            "value":   str(col_b.value).strip(),
        })
    return cells


# ── translation API ───────────────────────────────────────────────────────────

def call_api(file_name, cells, retries=3):
    """
    Send all cells for one file in a single batched API call.
    Returns (list[dict], stop_reason) on success, (None, error_str) on failure.
    """
    payload  = json.dumps(
        [{"row_idx": c["row_idx"], "label": c["label"], "value": c["value"]}
         for c in cells],
        ensure_ascii=False,
    )
    user_msg = f"File: {file_name!r}\n\n{payload}"
    last_err = None

    for attempt in range(retries):
        try:
            resp = _client.messages.create(
                model=MODEL,
                max_tokens=8192,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_msg}],
            )
            raw = resp.content[0].text.strip()
            # Strip markdown fences if model added them despite instructions
            if raw.startswith("```"):
                raw = "\n".join(raw.split("\n")[1:])
                raw = raw.rsplit("```", 1)[0].strip()
            parsed = json.loads(raw)
            if not isinstance(parsed, list):
                raise ValueError(f"Expected list, got {type(parsed).__name__}")
            return parsed, resp.stop_reason
        except Exception as e:
            last_err = e
            if attempt < retries - 1:
                time.sleep(2 ** attempt)

    return None, str(last_err)


# ── per-file pipeline ─────────────────────────────────────────────────────────

def process_file(fiche, write_drive=False):
    service   = get_service()
    file_id   = fiche["english_drive_id"]
    name      = fiche["english_name"]

    try:
        buf = download_bytes(service, file_id)
        wb  = openpyxl.load_workbook(buf)
    except Exception as e:
        return {"fiche": fiche, "error": f"download: {e}",
                "cells": [], "api_error": None, "stop_reason": None}

    ws    = wb.active
    cells = collect_cells(ws)

    if not cells:
        return {"fiche": fiche, "error": None,
                "cells": [], "api_error": None, "stop_reason": None}

    api_result, stop_reason = call_api(name, cells)

    if api_result is None:
        return {"fiche": fiche, "error": None, "cells": cells,
                "api_error": str(stop_reason), "stop_reason": None}

    by_row = {r["row_idx"]: r for r in api_result
              if isinstance(r, dict) and "row_idx" in r}

    cell_results = []
    for c in cells:
        ri   = c["row_idx"]
        item = by_row.get(ri)
        if item:
            translation = item.get("translation", c["value"])
            status      = item.get("status", "preserved_low_confidence")
        else:
            # Row missing from API response — safe default
            translation = c["value"]
            status      = "preserved_low_confidence"
        cell_results.append({
            "row_idx":     ri,
            "label":       c["label"],
            "original":    c["value"],
            "translation": translation,
            "status":      status,
        })

    # Write col C (openpyxl cell() creates the cell if it doesn't exist)
    for cr in cell_results:
        ws.cell(row=cr["row_idx"] + 1, column=3).value = cr["translation"]

    if write_drive:
        out_buf = io.BytesIO()
        wb.save(out_buf)
        out_buf.seek(0)
        media = MediaIoBaseUpload(
            out_buf,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resumable=False,
        )
        service.files().update(fileId=file_id, media_body=media).execute()
    else:
        LOCAL_DIR.mkdir(exist_ok=True)
        wb.save(LOCAL_DIR / name)

    return {"fiche": fiche, "error": None, "cells": cell_results,
            "api_error": None, "stop_reason": stop_reason}


# ── scope ──────────────────────────────────────────────────────────────────────

def get_targets():
    return [
        f for f in FICHES
        if f["subfolder"] in FICHES_SUBFOLDERS
        and f["english_name"] not in EXCLUDED_NAMES
        and f["english_name"].endswith(".xlsx")
    ]


# ── abbreviation scan ─────────────────────────────────────────────────────────

def scan_abbrevs(results):
    """
    Collect all 2-4 letter ALL-CAPS tokens from every col-B value.
    Returns {token: {"count": int, "examples": [...], "statuses": {status: count}}}.
    """
    found = {}
    for r in results:
        if r.get("error") or r.get("api_error"):
            continue
        fname = r["fiche"]["english_name"]
        for cr in r["cells"]:
            for m in _CAPS_RE.finditer(cr["original"]):
                tok = m.group(1)
                entry = found.setdefault(tok, {"count": 0, "examples": [], "statuses": {}})
                entry["count"] += 1
                s = cr.get("status", "?")
                entry["statuses"][s] = entry["statuses"].get(s, 0) + 1
                if len(entry["examples"]) < 3:
                    entry["examples"].append({
                        "file":    fname,
                        "row":     cr["row_idx"],
                        "label":   cr["label"],
                        "context": cr["original"][:100],
                    })
    return found


# ── report writers ────────────────────────────────────────────────────────────

def write_dry_report(results):
    low_conf   = []
    api_errors = []
    per_file   = []
    total_cells = translated = pres_name = pres_low = no_chg = 0

    for r in results:
        fname = r["fiche"]["english_name"]
        sub   = r["fiche"]["subfolder"]

        if r.get("error"):
            api_errors.append({"file": fname, "subfolder": sub, "error": r["error"]})
            per_file.append({"file": fname, "subfolder": sub,
                             "total_cells": 0, "translated": 0,
                             "preserved_code_or_name": 0, "preserved_low_confidence": 0,
                             "no_change": 0, "error": r["error"], "api_error": None})
            continue

        if r.get("api_error"):
            api_errors.append({"file": fname, "subfolder": sub, "error": r["api_error"]})
            n = len(r["cells"])
            per_file.append({"file": fname, "subfolder": sub,
                             "total_cells": n, "translated": 0,
                             "preserved_code_or_name": 0, "preserved_low_confidence": n,
                             "no_change": 0, "error": None, "api_error": r["api_error"]})
            continue

        ft = fp = fl = fn = 0
        for cr in r["cells"]:
            s = cr.get("status", "preserved_low_confidence")
            if s == "translated":               ft += 1
            elif s == "preserved_code_or_name": fp += 1
            elif s == "preserved_low_confidence": fl += 1
            elif s == "no_change":              fn += 1
            if s == "preserved_low_confidence":
                low_conf.append({
                    "file":     fname,
                    "subfolder": sub,
                    "row_idx":  cr["row_idx"],
                    "label":    cr["label"],
                    "original": cr["original"],
                })

        fc = ft + fp + fl + fn
        total_cells += fc
        translated  += ft
        pres_name   += fp
        pres_low    += fl
        no_chg      += fn
        per_file.append({
            "file": fname, "subfolder": sub,
            "total_cells": fc, "translated": ft,
            "preserved_code_or_name": fp,
            "preserved_low_confidence": fl,
            "no_change": fn,
            "error": None, "api_error": None,
        })

    abbrevs = scan_abbrevs(results)
    # Sort abbreviations by count descending for readability
    abbrevs_sorted = dict(
        sorted(abbrevs.items(), key=lambda x: -x[1]["count"])
    )

    with open(VALUE_DRY_REPORT, "w", encoding="utf-8") as fh:
        fh.write("# Generated by translate_values.py dry run\n")
        fh.write(f"STATS = {{\n")
        fh.write(f"    'total_files':             {len(results)},\n")
        fh.write(f"    'total_cells':             {total_cells},\n")
        fh.write(f"    'translated':              {translated},\n")
        fh.write(f"    'preserved_code_or_name':  {pres_name},\n")
        fh.write(f"    'preserved_low_confidence': {pres_low},\n")
        fh.write(f"    'no_change':               {no_chg},\n")
        fh.write(f"    'api_errors':              {len(api_errors)},\n")
        fh.write(f"}}\n\n")
        fh.write(f"LOW_CONFIDENCE = {low_conf!r}\n\n")
        fh.write(f"API_ERRORS = {api_errors!r}\n\n")
        fh.write(f"PER_FILE = {per_file!r}\n\n")
        fh.write(f"ABBREVS_FOUND = {abbrevs_sorted!r}\n")


def write_translation_log(results):
    entries = []
    for r in results:
        if r.get("error") or r.get("api_error") or not r["cells"]:
            continue
        changed = [cr for cr in r["cells"] if cr["status"] == "translated"]
        if not changed:
            continue
        f = r["fiche"]
        entries.append({
            "english_drive_id": f["english_drive_id"],
            "english_name":     f["english_name"],
            "subfolder":        f["subfolder"],
            "cells": changed,
        })
    with open(VALUE_TRANS_LOG, "w", encoding="utf-8") as fh:
        fh.write(f"VALUE_TRANSLATION_LOG = {entries!r}\n")


# ── main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true",
                        help="Write col C back to Drive (default: local dry run)")
    args = parser.parse_args()

    targets   = get_targets()
    mode      = "WRITE" if args.write else "DRY RUN"
    write_drv = args.write
    print(f"{mode}: {len(targets)} files")

    results = [None] * len(targets)
    done    = 0

    with concurrent.futures.ThreadPoolExecutor(max_workers=MAX_WORKERS) as pool:
        future_to_idx = {
            pool.submit(process_file, t, write_drv): i
            for i, t in enumerate(targets)
        }
        for future in concurrent.futures.as_completed(future_to_idx):
            idx = future_to_idx[future]
            results[idx] = future.result()
            done += 1
            if done % 50 == 0:
                print(f"  {done}/{len(targets)} done...")

    # ── collate ──
    dl_errors  = [r for r in results if r.get("error")]
    api_errors = [r for r in results if r.get("api_error")]
    good       = [r for r in results if not r.get("error") and not r.get("api_error")]
    all_cells  = [cr for r in good for cr in r["cells"]]

    total      = len(all_cells)
    translated = sum(1 for cr in all_cells if cr["status"] == "translated")
    pres_name  = sum(1 for cr in all_cells if cr["status"] == "preserved_code_or_name")
    pres_low   = sum(1 for cr in all_cells if cr["status"] == "preserved_low_confidence")
    no_chg     = sum(1 for cr in all_cells if cr["status"] == "no_change")

    print(f"\n{'='*60}")
    print(f"Files processed:               {len(results)}")
    print(f"Download/parse errors:         {len(dl_errors)}")
    print(f"API errors:                    {len(api_errors)}")
    print(f"Total col-B cells:             {total}")
    print(f"  translated:                  {translated}")
    print(f"  preserved_code_or_name:      {pres_name}")
    print(f"  preserved_low_confidence:    {pres_low}")
    print(f"  no_change (already English): {no_chg}")

    if dl_errors:
        print(f"\nDOWNLOAD/PARSE ERRORS:")
        for r in dl_errors:
            print(f"  {r['fiche']['english_name']!r}: {r['error']}")

    if api_errors:
        print(f"\nAPI ERRORS:")
        for r in api_errors:
            print(f"  {r['fiche']['english_name']!r}: {r['api_error']}")

    if not args.write:
        write_dry_report(results)
        print(f"\nFull detail  → value_dry_report.py")
        print(f"Local copies → dryrun_col_c/  ({len([r for r in results if not r.get('error') and not r.get('api_error') and r['cells']])} files written)")
        print("Run with --write to push col C to Drive.")
    else:
        write_translation_log(results)
        print(f"\nAudit trail → value_translation_log.py")


if __name__ == "__main__":
    main()
